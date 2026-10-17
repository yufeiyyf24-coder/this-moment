from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
src=Path('outputs/此刻_一周模拟记录.docx');d=Document(src)
original=[p.text for p in d.paragraphs];detail=original[original.index('2026年10月9日 星期五'):]
heading=next(p for p in d.paragraphs if p.text=='一周概览');heading.text='一周症状统计'
remove=False
for p in list(d.paragraphs):
 if p.text.startswith('疲惫：'):remove=True
 if p.text.startswith('未列出的标签'):
  p.text='平静、出门、社交、运动等其他记录保留在后文。未评分不代表零程度，未记录不代表未发生。程度仅供个人对照，不跨标签加总。'
  remove=False
 if remove:p._element.getparent().remove(p._element)
anchor=next(p for p in d.paragraphs if p.text.startswith('共14条记录'))
anchor.text='7天 · 14条记录。同一条记录中的同一症状计1次，按记录次数从多到少排列。'
symptoms=[('疲惫',4,4),('注意力差',3,3),('启动困难',2,2),('焦虑',2,2),('乏力',2,2),('胸闷',2,2),('头痛',1,1),('烦躁',1,1),('低落',1,1),('想哭',1,1),('说不清的难受',1,1),('头晕',1,1)]
t=d.add_table(rows=1,cols=3);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
widths=[3.45,1.6,1.6]
for c,w in zip(t.columns,widths):c.width=Inches(w)
for c,text in zip(t.rows[0].cells,['症状或不适','记录次数','出现天数']):c.text=text
for name,count,days in symptoms:
 cells=t.add_row().cells
 for c,text in zip(cells,[name,f'{count} 次',f'{days} 天']):c.text=text
for ri,row in enumerate(t.rows):
 trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
 if ri==0:
  rep=OxmlElement('w:tblHeader');trpr.append(rep)
 for ci,c in enumerate(row.cells):
  c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  pr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'354B32' if ri==0 else ('F0F4EA' if ri%2 else 'FFFFFF'));pr.append(sh)
  margins=OxmlElement('w:tcMar')
  for side,value in [('top',65),('bottom',65),('left',140),('right',140)]:
   node=OxmlElement('w:'+side);node.set(qn('w:w'),str(value));node.set(qn('w:type'),'dxa');margins.append(node)
  pr.append(margins)
  p=c.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.05
  for r in p.runs:
   r.font.size=Pt(12);r.bold=ri==0 or ci==1;r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000');r.font.name='Microsoft YaHei';r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
anchor._p.addnext(t._tbl)
# Preserve subsequent detailed entries exactly.
now=[p.text for p in d.paragraphs];assert now[now.index('2026年10月9日 星期五'):]==detail
assert len(t.rows)==13
out=Path('outputs/此刻_一周模拟记录_症状统计表版.docx');d.save(out)
print('PASS: 12 symptoms, prominent table, unchanged daily details. File:',out.name)
