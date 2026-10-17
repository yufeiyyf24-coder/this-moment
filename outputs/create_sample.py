from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from collections import Counter,defaultdict
from pathlib import Path
D=Document();sec=D.sections[0];sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7);sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=D.styles[name];st.font.name='Microsoft YaHei';st._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei');st.font.color.rgb=RGBColor(0,0,0)
D.styles['Normal'].font.size=Pt(11);D.styles['Normal'].paragraph_format.space_after=Pt(5);D.styles['Normal'].paragraph_format.line_spacing=1.25
D.styles['Title'].font.size=Pt(23)
D.styles['Heading 1'].font.size=Pt(15);D.styles['Heading 1'].paragraph_format.space_before=Pt(14)
D.styles['Heading 2'].font.size=Pt(11);D.styles['Heading 2'].paragraph_format.space_before=Pt(9)
rows=[
('10月9日 星期五', [('09:20',[('疲惫',2,None),('注意力差',2,None)],'昨晚睡得晚。早上读东西时需要重看几遍。'),('18:40',[('平静',None,None),('出门',None,None)],'去附近走了十几分钟，回来觉得安稳一些。')]),
('10月10日 星期六',[('11:10',[('启动困难',3,None),('焦虑',2,None)],'想收拾房间，东西准备好了，还是坐了很久才开始。'),('21:30',[('疲惫',3,None),('头痛',2,'头部')],'下午做了一些家务，晚上想早点休息。')]),
('10月11日 星期日',[('15:05',[('胸闷',2,'胸部'),('焦虑',3,None)],'出门前反复检查东西，胸口有点闷。坐了一会儿才出门。'),('20:15',[('社交',None,None),('平静',None,None)],'和朋友吃饭。刚开始话少，后来聊得比较轻松。')]),
('10月12日 星期一',[('08:50',[('乏力',3,'全身'),('疲惫',3,None)],'起床后手脚沉，洗漱到一半坐下来休息了一会儿。'),('16:20',[('注意力差',3,None),('烦躁',2,None)],'整理资料时不断切换页面，有人说话时更容易分心。')]),
('10月13日 星期二',[('13:10',[('低落',3,None),('想哭',2,None)],'午饭后不太想回应消息，也说不清具体原因。'),('22:00',[('说不清的难受',None,None)],'不想选程度。记下这一刻，明天再看看。')]),
('10月14日 星期三',[('10:35',[('启动困难',2,None),('注意力差',2,None)],'把事情拆成一小步后开始了，做了大约二十分钟。'),('19:10',[('运动',None,None),('平静',None,None)],'在家伸展了十分钟，身体比下午放松。')]),
('10月15日 星期四',[('09:15',[('疲惫',2,None),('头晕',1,'头部')],'起身时有点晕，放慢动作后继续洗漱。'),('14:30',[('乏力',3,'全身'),('胸闷',None,'胸部')],'下午出门后觉得很累，坐了一会儿。这次胸闷没有评分。')])]
D.add_paragraph('此刻一周记录示例','Title')
D.add_paragraph('2026年10月9日至10月15日')
p=D.add_paragraph();r=p.add_run('虚构示例数据');r.bold=True;p.add_run('  仅用于查看导出样式，不代表你的真实经历。')
D.add_paragraph('这份记录按日期和时间保留标签、程度、部位及原话，供就诊时查看，也可直接在手机 Word 或 WPS 中编辑。')
D.add_heading('一周概览',1)
D.add_paragraph('共14条记录，记录覆盖7天。以下统计按每次记录计数，同一条记录中同一标签计1次。')
counts=Counter();days=defaultdict(set)
for date,entries in rows:
 for time,tags,note in entries:
  for tag,level,part in tags:counts[tag]+=1;days[tag].add(date)
for name in ['疲惫','注意力差','平静','启动困难','焦虑','乏力','胸闷']:
 D.add_paragraph(f'{name}：{counts[name]}次，出现在{len(days[name])}天。')
D.add_paragraph('未列出的标签各记录1次。未评分不代表零程度，未记录不代表未发生。1至5档仅用于本人回看同一标签，不跨标签加总，不作为临床量表或诊断结论。')
D.add_heading('想在就诊时补充的内容',1)
D.add_paragraph('可在这里补写想问的问题、近期变化，或需要向医生解释的背景。')
# Two days per page, with the last three days on the final page.
for idx,(date,entries) in enumerate(rows):
 if idx in (0,2,4):D.add_page_break()
 D.add_heading('2026年'+date,1)
 for time,tags,note in entries:
  D.add_heading(time,2)
  for tag,level,part in tags:
   p=D.add_paragraph();p.paragraph_format.keep_with_next=True
   r=p.add_run(tag);r.bold=True
   if part:p.add_run('  |  '+part)
   p.add_run('  |  '+(f'程度 {level}/5' if level else '未评分' if tag in ['胸闷'] else '仅记录'))
  p=D.add_paragraph('补充：'+note);p.paragraph_format.space_after=Pt(12)
# Repeating footer makes the simulated nature unambiguous even on separated pages.
p=sec.footer.paragraphs[0];p.alignment=2
r=p.add_run('此刻  虚构示例  ·  ');r.font.size=Pt(9)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
D.core_properties.title='此刻一周记录示例';D.core_properties.author='';D.core_properties.subject='虚构的一周记录导出样式'
out=Path('outputs/此刻_一周模拟记录.docx');D.save(out);print(out.resolve())
