from pathlib import Path
import ast
from collections import defaultdict
from docx import Document
from docx.shared import Pt
src=Path('outputs/此刻_一周模拟记录_症状统计表版.docx');d=Document(src)
code=ast.parse(Path('outputs/create_sample.py').read_text(encoding='utf-8-sig'))
rows=next(ast.literal_eval(n.value) for n in code.body if isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='rows' for t in n.targets))
days=defaultdict(set);scores=defaultdict(list)
for day,entries in rows:
 for time,tags,note in entries:
  for name,score,part in tags:
   days[name].add(day)
   if score is not None:scores[name].append(score)
original=[p.text for p in d.paragraphs];detail=original[original.index('2026年10月9日 星期五'):]
t=d.tables[0]
# Keep existing row order among equal day counts.
items=[r.cells[0].text for r in t.rows[1:]];items.sort(key=lambda name:-len(days[name]))
for c,text in zip(t.rows[0].cells,['症状或不适','出现天数','程度均分']):
 c.paragraphs[0].runs[0].text=text
for row,name in zip(t.rows[1:],items):
 values=[name,f'{len(days[name])} 天',f'{sum(scores[name])/len(scores[name]):.1f} / 5' if scores[name] else '未评分']
 for col,(cell,value) in enumerate(zip(row.cells,values)):
  runs=cell.paragraphs[0].runs;runs[0].text=value
  for r in runs[1:]:r.text=''
  runs[0].bold=col==2;runs[0].font.size=Pt(12)
for p in d.paragraphs:
 if p.text.startswith('7天 · 14条记录。'):p.text='7天 · 14条记录。按症状出现天数从多到少排列。'
 if p.text.startswith('平静、出门、社交、运动等其他记录'):
  p.text='出现天数：当天记录过该症状即计1天。程度均分：仅对同一症状的已评分记录取平均，保留1位小数；未评分不按0计算。胸闷共记录2天，仅1条有评分。平静及活动记录见后文。未记录不代表未发生，程度仅供个人对照。'
  p.paragraph_format.space_after=Pt(5)
  for r in p.runs:r.font.size=Pt(10)
assert [p.text for p in d.paragraphs][original.index('2026年10月9日 星期五'):]==detail
assert scores['疲惫']==[2,3,3,2]
assert sum(scores['胸闷'])/len(scores['胸闷'])==2
assert len(days['胸闷'])==2
out=Path('outputs/此刻_一周模拟记录_天数与程度均分.docx');d.save(out)
check=Document(out)
assert check.tables[0].rows[1].cells[2].text=='2.5 / 5'
print('PASS: day counts, averages exclude unrated records, original details preserved.')
